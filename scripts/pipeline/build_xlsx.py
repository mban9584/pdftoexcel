"""Turn cells/m{0..4}.json into C:\\Users\\sample\\Desktop\\out.xlsx (A4, original layout).

The page is tiled by an extended grid XG = [0]+xs+[W] (same for rows) so that every table
cell, margin caption and photo maps onto cell ranges. Column widths: px = pt*4/3,
width = (px-5)/7. Row heights: the pt band itself.
"""
import json
import os
import re
import sys

PDF_PATH = os.environ.get('PDFTOEXCEL_PDF', os.path.expanduser('~/Desktop/sample.pdf'))
WORK = os.environ.get('PDFTOEXCEL_WORK', os.path.expanduser('~/pdftoexcel-work'))
OUT_PATH = os.path.join(os.path.expanduser('~/Desktop'), 'out.xlsx')
from bisect import bisect_right

import numpy as np
import pypdfium2 as pdfium
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

PDF = PDF_PATH
DIR = os.path.join(WORK, 'cells')
IMGDIR = os.path.join(WORK, 'imgs')
OUT = OUT_PATH
EMU_PT = 12700.0
SCALE = 300 / 72.0            # px per pt for photo crops
INK_SCALE = 4.0               # px per pt for ink measurement
FONT = '宋体'
BODY_PT = 12.0                # pages 2-4, calibrated from rendered digit ink
CJK_INK_RATIO = 0.917         # p1: 12pt CJK glyph measures 11.0pt of ink
SLASH_W = (2.5, 9.0)          # 无货 "/" glyph bbox in the original
SLASH_H = (7.0, 14.0)
THIN = Side(style='thin', color='FF000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLOR = {'red': 'FFFF0000', 'blue': 'FF0000FF', 'black': 'FF000000', 'white': 'FFFFFFFF'}

doc = pdfium.PdfDocument(PDF)
_cache = {}


def grey(page):
    if page not in _cache:
        _cache[page] = np.asarray(doc[page].render(scale=INK_SCALE).to_pil().convert('L'))
    return _cache[page]


def colour(page):
    key = ('rgb', page)
    if key not in _cache:
        _cache[key] = doc[page].render(scale=SCALE).to_pil()
    return _cache[key]


def ink_box(page, x0, y0, x1, y1, pad=1.2):
    """Ink bbox in pt inside a region, ignoring table lines; None when blank."""
    a = grey(page)
    sub = a[int((y0 + pad) * INK_SCALE):int((y1 - pad) * INK_SCALE),
            int((x0 + pad) * INK_SCALE):int((x1 - pad) * INK_SCALE)]
    if sub.size == 0:
        return None
    dark = sub < 150
    dark = dark & (dark.mean(axis=1) < 0.85)[:, None] & (dark.mean(axis=0) < 0.85)[None, :]
    rows = np.where(dark.any(axis=1))[0]
    cols = np.where(dark.any(axis=0))[0]
    if len(rows) == 0:
        return None
    return (x0 + pad + cols[0] / INK_SCALE, y0 + pad + rows[0] / INK_SCALE,
            x0 + pad + (cols[-1] + 1) / INK_SCALE, y0 + pad + (rows[-1] + 1) / INK_SCALE)


def span(a, b, G, n):
    """Extended-grid indices (0-based, inclusive) covered by the point range [a, b)."""
    k0 = max(0, min(bisect_right(G, a + 0.4) - 1, n - 1))
    k1 = max(k0, min(bisect_right(G, b - 0.4) - 1, n - 1))
    return k0, k1


def snap_grid(vals, total):
    """Boundaries in pt, snapped to the pixel grid.

    Excel realises a column as round(width*7+5) px, so snapping each *boundary* (not each
    band) keeps every edge within half a pixel of the PDF instead of accumulating.
    """
    b = [0] + [int(round(v * 4 / 3)) for v in list(vals) + [total]]
    for k in range(1, len(b)):
        b[k] = max(b[k], b[k - 1] + 1)
    px = [b[k + 1] - b[k] for k in range(len(b) - 1)]
    return [x * 0.75 for x in b], px


def sheet_name(m, page, used):
    if page == 0:
        base = '封面'
    else:
        base = (max(m['free'], key=lambda f: f['size'])['text'].strip() if m['free'] else 'p%d' % page)
    base = re.sub(r'[\\/*?:\[\]]', '', base)[:28] or 'p%d' % page
    name, k = base, 1
    while name in used:
        k += 1
        name = '%s（%s）' % (base, '二三四五六'[k - 2])
    used.add(name)
    return name


os.makedirs(IMGDIR, exist_ok=True)
wb = Workbook()
wb.remove(wb.active)
used, report = set(), []

for page in ([int(x) for x in sys.argv[1:]] or [0, 1, 2, 3, 4]):
    m = json.load(open(os.path.join(DIR, 'm%d.json' % page), encoding='utf-8'))
    W, H = float(m['w']), float(m['h'])
    XG, XPX = snap_grid([float(x) for x in m['xs']], W)
    YG = [0.0] + [float(y) for y in m['ys']] + [H]
    ncol, nrow = len(XG) - 1, len(YG) - 1
    ws = wb.create_sheet(sheet_name(m, page, used))
    ws.sheet_view.showGridLines = False
    for k in range(ncol):
        if XPX[k] <= 5:
            print('   warn p%d col%d width %dpx -> clamped to 0.06' % (page, k + 1, XPX[k]))
        ws.column_dimensions[get_column_letter(k + 1)].width = max((XPX[k] - 5) / 7, 0.06)
    for k in range(nrow):
        ws.row_dimensions[k + 1].height = max(YG[k + 1] - YG[k], 1.0)

    # table cells -> extended-grid ranges (original index i becomes column i+1)
    ranges = []
    taken = [[False] * nrow for _ in range(ncol)]
    for c in m['cells']:
        i0, j0, i1, j1 = c['cell']
        a0, a1, b0, b1 = i0 + 2, i1 + 1, j0 + 2, j1 + 1
        ranges.append((c, a0, a1, b0, b1, (XG[i0 + 1], YG[j0 + 1], XG[i1 + 1], YG[j1 + 1])))
        for x in range(a0, min(a1, ncol) + 1):
            for y in range(b0, min(b1, nrow) + 1):
                taken[x - 1][y - 1] = True

    def put(a0, a1, b0, b1, value, font, align, border):
        if a1 > a0 or b1 > b0:
            ws.merge_cells(start_row=b0, start_column=a0, end_row=b1, end_column=a1)
        for x in range(a0, min(a1, ncol) + 1):
            for y in range(b0, min(b1, nrow) + 1):
                cell = ws.cell(row=y, column=x)
                cell.font = font
                if border:
                    cell.border = BOX
        head = ws.cell(row=b0, column=a0)
        head.alignment = align
        if value is not None:
            head.value = value
        return head

    n_text = n_slash = 0
    for c, a0, a1, b0, b1, box in ranges:
        text = c['text'].strip()
        if not text:
            ink = ink_box(page, *box)
            if ink and SLASH_W[0] <= ink[2] - ink[0] <= SLASH_W[1] and SLASH_H[0] <= ink[3] - ink[1] <= SLASH_H[1]:
                text, n_slash = '/', n_slash + 1
        size = c['size'] if page <= 1 else BODY_PT
        has_img = False
        for im in m['images']:
            dx = min(box[2], im[2]) - max(box[0], im[0])
            dy = min(box[3], im[3]) - max(box[1], im[1])
            if dx > 0 and dy > 0 and dx * dy > 0.2 * max(1.0, (im[2] - im[0]) * (im[3] - im[1])):
                has_img = True
                break
        align = Alignment(horizontal='center', vertical='bottom' if (has_img and b1 > b0) else 'center',
                          wrap_text=True)
        font = Font(name=FONT, size=round(size, 1), color=COLOR.get(c['color'], 'FF000000'))
        value = text or None
        if text and re.fullmatch(r'\d+(\.\d+)?', text):
            dp = len(text.split('.')[1]) if '.' in text else 0
            value = int(text) if dp == 0 else float(text)
            head = put(a0, a1, b0, b1, value, font, align, True)
            head.number_format = '0' if dp == 0 else '0.' + '0' * dp
        else:
            put(a0, a1, b0, b1, value, font, align, True)
        if text:
            n_text += 1

    n_free = 0
    for f in m['free']:
        if not f['text'].strip():
            continue
        if page == 0:
            a0, a1, b0, b1 = f['i0'] + 2, f['i1'] + 1, f['j0'] + 2, f['j1'] + 1
        else:
            (a0, a1), (b0, b1) = span(f['x0'], f['x1'], XG, ncol), span(f['y0'], f['y1'], YG, nrow)
            a0, a1, b0, b1 = a0 + 1, a1 + 1, b0 + 1, b1 + 1
        a1 = max(a1, a0)
        b1 = max(b1, b0)
        if any(taken[x - 1][y - 1] for x in range(a0, min(a1, ncol) + 1) for y in range(b0, min(b1, nrow) + 1)):
            a1, b1 = a0, b0          # occupied elsewhere: let it overflow instead of merging
        size = f['size'] if page <= 1 else max(round((f['y1'] - f['y0'] - 1.2) / CJK_INK_RATIO, 1), 8.0)
        font = Font(name=FONT, size=round(size, 1), color=COLOR.get(f['color'], 'FF000000'))
        put(min(a0, ncol), min(a1, ncol), min(b0, nrow), min(b1, nrow), f['text'].strip(), font,
            Alignment(horizontal='center', vertical='center', wrap_text=True), False)
        n_free += 1

    n_img = 0
    for n, im in enumerate(m['images']):
        x0, y0, x1, y1 = im
        img = colour(page).crop(tuple(int(v * SCALE) for v in (x0, y0, x1, y1)))
        if img.size[0] < 2 or img.size[1] < 2:
            continue
        path = os.path.join(IMGDIR, 'p%d_%d.png' % (page, n))
        img.save(path)
        a0 = span(x0, x0 + 1, XG, ncol)[0]
        b0 = span(y0, y0 + 1, YG, nrow)[0]
        pic = XImage(path)
        pic.anchor = OneCellAnchor(
            _from=AnchorMarker(col=a0, colOff=int((x0 - XG[a0]) * EMU_PT),
                               row=b0, rowOff=int((y0 - YG[b0]) * EMU_PT)),
            ext=XDRPositiveSize2D(int((x1 - x0) * EMU_PT), int((y1 - y0) * EMU_PT)))
        ws.add_image(pic)
        n_img += 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.2, bottom=0.2, header=0.05, footer=0.05)
    ws.print_area = 'A1:%s%d' % (get_column_letter(ncol), nrow)
    report.append((ws.title, page, ncol, nrow, len(ranges), n_text, n_slash, n_free, n_img))

wb.save(OUT)
for r in report:
    print('sheet %-24s p%d cols=%-3d rows=%-3d cells=%-4d text=%-4d slash=%d free=%-3d imgs=%d' % r)
print('saved %s (%d bytes)' % (OUT, os.path.getsize(OUT)))
