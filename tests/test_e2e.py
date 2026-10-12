"""probe -> model -> xlsx -> check on a generated PDF, through the real CLI.

The fixture is drawn with reportlab (requirements-dev.txt): three portrait price-list
pages with CJK text and a landscape page, all with a text layer, so no OCR is needed.
"""
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CLI = str(ROOT / "scripts" / "pdftoexcel.py")

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
except ImportError:          # pragma: no cover - CI installs it
    canvas = None


def _table(c, x0, y0, rows, cw, rh):
    cols = len(rows[0])
    for r in range(len(rows) + 1):
        c.line(x0, y0 - r * rh, x0 + cols * cw, y0 - r * rh)
    for k in range(cols + 1):
        c.line(x0 + k * cw, y0, x0 + k * cw, y0 - len(rows) * rh)
    for r, row in enumerate(rows):
        for k, t in enumerate(row):
            c.drawCentredString(x0 + k * cw + cw / 2, y0 - r * rh - rh * 0.65, t)


def make_pdf(path):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(path, pagesize=A4)
    for p in range(3):
        c.setFont("STSong-Light", 16)
        c.drawString(150, 790, "2026 年度产品报价单 第%d页" % (p + 1))
        c.setFont("STSong-Light", 10)
        rows = [["名称", "规格", "单价"]] + [["产品%d" % i, "25×20", "%d.50" % (i * 10)] for i in range(1, 22)]
        rows[3][2] = "020"          # leading zero: must stay text
        _table(c, 60, 760, rows, 150, 20)
        c.showPage()
    c.setPageSize(landscape(A4))
    c.setFont("STSong-Light", 10)
    _table(c, 40, 560, [["横%d" % k for k in range(6)]] + [[str(r * k) for k in range(6)] for r in range(1, 22)],
           120, 20)
    c.showPage()
    c.save()


@unittest.skipIf(canvas is None, "reportlab not installed (pip install -r requirements-dev.txt)")
class EndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        d = Path(cls.tmp.name)
        cls.pdf = str(d / "报价单 test.pdf")      # CJK + space in the path on purpose
        cls.env = dict(os.environ, PDFTOEXCEL_WORK=str(d / "work"), PYTHONIOENCODING="utf-8")
        make_pdf(cls.pdf)
        cls.run_cli("probe", "--pdf", cls.pdf)
        cls.run_cli("model", "--pdf", cls.pdf)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    @classmethod
    def run_cli(cls, *args):
        p = subprocess.run([sys.executable, CLI, *args], env=cls.env, capture_output=True,
                           text=True, encoding="utf-8")
        if p.returncode:
            raise AssertionError("%s failed:\n%s\n%s" % (args[0], p.stdout, p.stderr))
        return p.stdout

    def xlsx(self, name, *extra):
        out = str(Path(self.tmp.name) / name)
        self.run_cli("xlsx", "--pdf", self.pdf, "--out", out, *extra)
        return out

    def test_full_build_checks_ok(self):
        out = self.xlsx("all.xlsx")
        self.assertTrue(self.run_cli("check", "--pdf", self.pdf, "--xlsx", out).rstrip().endswith("OK"))

    def test_check_maps_sheets_to_pages_after_partial_build(self):
        out = self.xlsx("subset.xlsx", "--pages", "2,3")
        report = self.run_cli("check", "--pdf", self.pdf, "--xlsx", out)
        self.assertNotIn("MISMATCH", report)
        self.assertTrue(report.rstrip().endswith("OK"), report)

    def test_layout_details(self):
        wb = load_workbook(self.xlsx("details.xlsx"))
        first, wide = wb.worksheets[0], wb.worksheets[3]
        self.assertNotEqual(first.title, "封面")                   # page 0 is a table, not a cover
        self.assertEqual(first.page_setup.orientation, "portrait")
        self.assertEqual(wide.page_setup.orientation, "landscape")
        cells = {c.value: c for row in first.iter_rows() for c in row if c.value is not None}
        self.assertEqual(cells[20.5].number_format, "0.00")       # "20.50" -> number, 2 dp kept
        self.assertIn("020", cells)                               # leading zero stays text

    def test_fontsize_runs(self):
        self.assertIn("ink height", self.run_cli("fontsize", "--pdf", self.pdf, "--page", "0",
                                                 "--rect", "60,40,360,60"))


if __name__ == "__main__":
    unittest.main()
